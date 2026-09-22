from __future__ import annotations

import json
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "LUMENFALL_Job_and_Skill_Progression.docx"
NODE = Path(r"C:\Users\USER\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")


def load_registry() -> dict:
    js = """
import('./lib/game/skills.ts').then(m => console.log(JSON.stringify({
  core: Object.values(m.CORE_JOBS),
  special: Object.values(m.SPECIALIZATIONS),
  adventurer: m.ADVENTURER_SKILLS,
  coreSkills: m.CORE_SKILLS,
  specialSkills: m.SPECIAL_SKILLS,
  passives: m.PASSIVES,
  passiveEffects: m.PASSIVE_EFFECTS,
})))
"""
    result = subprocess.run(
        [str(NODE), "--experimental-strip-types", "-e", js],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


DATA = load_registry()

CORE_ORDER = ["warrior", "rogue", "hunter", "wizard", "acolyte"]
SPECIAL_ORDER = {
    "warrior": ["gatotkaca", "garda"],
    "rogue": ["caroq", "anom"],
    "hunter": ["srikandi", "jagawana"],
    "wizard": ["resi", "pujangga"],
    "acolyte": ["pandita", "bajra"],
}
CORE_NAMES = {item["id"]: item["name"] for item in DATA["core"]}
SPECIALS = {item["name"].lower(): item for item in DATA["special"]}

WEAPON_LABELS = {
    "mace": "Gada",
    "dual_dagger": "Dual dagger",
    "bow": "Busur",
    "staff": "Staff",
    "sword_shield": "Pedang dan perisai",
    "knuckle": "Knuckle dua tangan",
    "sword_dagger": "Pedang dan dagger",
    "bow_trap": "Busur perangkap",
    "wand": "Wand",
    "talisman": "Talisman",
    "relic": "Relic",
    "holy_knuckle": "Holy knuckle",
}
EFFECT_LABELS = {
    "damage": "Damage",
    "dash_damage": "Dash + damage",
    "rapid_damage": "Rapid damage",
    "aoe_damage": "Damage area",
    "heal": "Heal",
    "buff": "Buff",
    "debuff": "Debuff",
    "poison": "Poison",
    "slow": "Slow",
    "root": "Root",
    "stun": "Stun",
    "mark": "Mark",
    "stealth": "Stealth",
    "barrier": "Barrier",
    "parry": "Parry",
    "execute": "Execute",
    "elemental": "Elemental",
    "chain": "Chain",
    "illusion": "Illusion",
    "ultimate": "Ultimate",
}
STATUS_LABELS = {
    "knockup": "Knock-up",
    "taunt": "Taunt",
    "defenseDown": "Defense Down",
    "critical": "Critical",
    "weakPoint": "Weak Point",
    "slow": "Slow",
    "root": "Root",
    "poison": "Poison",
    "curse": "Curse",
    "stagger": "Stagger",
    "barrier": "Barrier",
}


def text(value: object) -> str:
    return str(value) if value is not None else "-"


def number(value: object) -> str:
    if value is None:
        return "-"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return f"{value:g}" if isinstance(value, float) else str(value)


def seconds(value: object) -> str:
    return f"{float(value):.1f}s".replace(".0s", "s")


def weapon_list(values: list[str]) -> str:
    return ", ".join(WEAPON_LABELS.get(value, value) for value in values) if values else "Semua senjata"


def set_cell_shading(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_run(run, size=9.5, bold=False, color="222222", italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=5, line=1.08):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line


def add_heading(doc, label, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(label)
    style_run(run, size=15 if level == 1 else 11.5, bold=True, color="000000")
    return paragraph


def add_body(doc, content, bold_lead=None):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, space_after=7, line=1.12)
    if bold_lead and content.startswith(bold_lead):
        style_run(paragraph.add_run(bold_lead), size=10.2, bold=True)
        style_run(paragraph.add_run(content[len(bold_lead):]), size=10.2)
    else:
        style_run(paragraph.add_run(content), size=10.2)
    return paragraph


def add_table(doc, headers, rows, widths, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_widths(table, widths)
    header = table.rows[0]
    repeat_header(header)
    for cell, label in zip(header.cells, headers):
        set_cell_shading(cell, "243B53")
        set_cell_margins(cell, top=110, bottom=110)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(paragraph, space_after=0, line=1.0)
        style_run(paragraph.add_run(label), size=font_size, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            set_cell_shading(cell, "F4F7FA" if row_index % 2 else "FFFFFF")
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            style_paragraph(paragraph, space_after=0, line=1.02)
            style_run(paragraph.add_run(text(value)), size=font_size)
    set_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def skill_row(skill):
    unlock = f"Slot {skill['slot']}\nLv. {skill['unlockLevel']}"
    name = f"{skill['name']}\n{skill['description']}"
    effect = EFFECT_LABELS.get(skill["effect"], skill["effect"])
    status = STATUS_LABELS.get(skill.get("statusEffect"), "") if skill.get("statusEffect") else ""
    if status:
        effect += f" · {status}"
    damage = ""
    if skill.get("baseDamage", 0):
        damage = f"Damage {number(skill['baseDamage'])} x {number(skill.get('damageCoefficient', 1))}"
    effect = f"{effect}\n{damage}" if damage else effect
    mana_cd = f"Mana {number(skill['manaCost'])}\nCD {seconds(skill['cooldown'])}"
    target = f"{skill['targetType']} · {number(skill['range'])}m"
    if skill.get("areaRadius", 0):
        target += f"\nArea {number(skill['areaRadius'])}m"
    if skill.get("weaponRequirement"):
        target += f"\n{weapon_list(skill['weaponRequirement'])}"
    return [unlock, name, effect, mana_cd, target]


def add_skill_table(doc, skills):
    return add_table(
        doc,
        ["Urutan dan buka", "Skill", "Fungsi", "Biaya dan cooldown", "Target dan syarat"],
        [skill_row(item) for item in sorted(skills, key=lambda x: x["slot"])],
        [0.72, 2.22, 1.42, 1.02, 1.62],
        font_size=8.3,
    )


def add_passive_table(doc, passive, effect, capstone=False):
    effect_text = ", ".join(f"{key} +{number(value)}" for key, value in effect.items()) if effect else "-"
    max_level = "1" if capstone else str(passive.get("maxLevel", 3))
    requirement = "Mastery Quest + passive utama Lv. 3" if capstone else f"Lv. {passive.get('unlockLevel', 25)}"
    add_table(
        doc,
        ["Passive", "Fungsi", "Bonus registry", "Level maksimum", "Syarat"],
        [[passive["name"], passive["description"], effect_text, max_level, requirement]],
        [1.35, 2.35, 1.65, 0.8, 0.85],
        font_size=8.4,
    )


def add_job_profile(doc, job, specializations):
    rows = [
        ["Peran", job["role"]],
        ["Senjata utama", WEAPON_LABELS.get(job["weapon"], job["weapon"])],
        ["HP dasar dan pertumbuhan", f"{number(job['hp'])} + {number(job['hpGrowth'])} per level"],
        ["Attack dasar dan pertumbuhan", f"{number(job['attack'])} + {number(job['attackGrowth'])} per level"],
        ["Basic cooldown dan jangkauan", f"{seconds(job['cooldown'])} · {number(job['range'])}m"],
        ["Passive dasar", job["passiveName"]],
        ["Turunan", " dan ".join(item["name"] for item in specializations)],
        ["Ringkasan", job["description"]],
    ]
    add_table(doc, ["Atribut", "Data Lumenfall"], rows, [2.0, 5.0], font_size=9.0)


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(paragraph, space_after=0, line=1.0)
    style_run(paragraph.add_run("Lumenfall · Katalog Job dan Skill · Data registry project"), size=8.5, color="667085")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    style_run(title.add_run("LUMENFALL Job dan Skill"), size=25, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    style_run(subtitle.add_run("Struktur progression Adventurer, Core Job, Special Job dan passive tree"), size=12.5, color="53657A")
    meta = doc.add_paragraph()
    style_run(meta.add_run(f"Katalog implementasi · diperbarui {date.today().strftime('%d %B %Y')}"), size=9.2, color="667085", italic=True)

    add_body(doc, "Dokumen ini menunjukkan seluruh jalur Job Lumenfall secara berurutan, dari Adventurer sampai Special Job dan Capstone. Setiap skill ditulis berdasarkan registry aktif project: nama, level buka, fungsi, biaya Mana, cooldown, target, dan syarat senjata bila ada.")
    add_body(doc, "Kesimpulan utama: Lumenfall memiliki 1 tahap awal, 5 Core Job, 10 Special Job, 64 active skill, dan 26 passive node. Skill Adventurer tetap tersedia setelah promosi; Core Job menambah empat skill job, sedangkan Special Job menambah empat skill spesialisasi.", bold_lead="Kesimpulan utama:")

    add_heading(doc, "Urutan progression", 1)
    progression = [
        ["1", "Adventurer", "Mulai permainan", "4 active skill + Tekad Petualang", "Level 1"],
        ["2", "Core Job", "Pilih 1 dari 5 jalur", "4 active skill + passive foundation", "Level 10; equipment harus dilepas"],
        ["3", "Special Job", "Pilih 1 dari 2 turunan Core Job", "4 active skill + passive spesialisasi", "Level 25; specialization harus sesuai Core Job"],
        ["4", "Mastery", "Pilih gaya Power, Control atau Utility", "Modifier pada active skill", "Level 40; memilih skill aktif"],
        ["5", "Capstone", "Puncak jalur Special Job", "Warisan passive dengan bonus ganda", "Level 50; Mastery + passive utama Lv. 3"],
    ]
    add_table(doc, ["Tahap", "Nama", "Arah", "Yang terbuka", "Syarat"], progression, [0.55, 1.05, 1.55, 2.2, 1.65], font_size=8.8)

    add_heading(doc, "Mastery active skill", 2)
    add_body(doc, "Setiap active skill memiliki tiga pilihan Mastery. Power meningkatkan damage menjadi 1.25 kali, Control meningkatkan durasi menjadi 1.45 kali, dan Utility mengurangi cooldown menjadi 0.78 kali. Mastery dipilih per skill setelah karakter mencapai level 40 dan masuk jalur Special Job.")

    add_heading(doc, "Peta Core Job", 1)
    core_rows = []
    for core_id in CORE_ORDER:
        job = next(item for item in DATA["core"] if item["id"] == core_id)
        specs = [next(item for item in DATA["special"] if item["name"].lower() == name.lower()) for name in SPECIAL_ORDER[core_id]]
        core_rows.append([job["name"], job["role"], WEAPON_LABELS.get(job["weapon"], job["weapon"]), " / ".join(item["name"] for item in specs), job["description"]])
    add_table(doc, ["Core Job", "Peran", "Senjata", "Turunan", "Identitas"], core_rows, [1.0, 1.35, 1.2, 1.3, 2.15], font_size=8.6)

    add_heading(doc, "Adventurer", 1)
    add_body(doc, "Semua karakter memulai sebagai Adventurer. Empat skill ini tersedia sebelum memilih Core Job dan tetap dapat digunakan setelah promosi. Tekad Petualang adalah passive awal yang dapat dinaikkan sampai level 3 mulai level 1.")
    add_skill_table(doc, DATA["adventurer"])
    add_passive_table(doc, {"name": "Tekad Petualang", "description": "Setiap level memberi +10 HP dan +1 Defense.", "maxLevel": 3, "unlockLevel": 1}, DATA["passiveEffects"]["adventurer-resolve"])

    core_skills = {core_id: [item for item in DATA["coreSkills"] if item["job"] == core_id] for core_id in CORE_ORDER}
    special_skills = {spec: [item for item in DATA["specialSkills"] if item["specialization"] == spec] for spec in DATA["passives"]}
    for index, core_id in enumerate(CORE_ORDER):
        job = next(item for item in DATA["core"] if item["id"] == core_id)
        specs = [next(item for item in DATA["special"] if item["name"].lower() == name.lower()) for name in SPECIAL_ORDER[core_id]]
        foundation_id = f"{core_id}-foundation"
        foundation_effect = DATA["passiveEffects"].get(foundation_id, {})
        add_heading(doc, f"{index + 1}. {job['name']}", 1)
        add_body(doc, job["description"] + f" Jalur ini memakai {WEAPON_LABELS.get(job['weapon'], job['weapon'])} dan membuka dua Special Job: " + " dan ".join(item["name"] for item in specs) + ".")
        add_job_profile(doc, job, specs)
        add_heading(doc, f"Skill inti {job['name']}", 2)
        add_skill_table(doc, core_skills[core_id])
        foundation = {"name": job["passiveName"], "description": f"Latihan dasar {job['name']} memperkuat atribut utama job.", "maxLevel": 3, "unlockLevel": 10}
        add_heading(doc, f"Passive foundation {job['name']}", 2)
        add_passive_table(doc, foundation, foundation_effect)
        for spec in specs:
            spec_key = spec["name"].lower()
            passive = DATA["passives"][spec_key]
            add_heading(doc, f"Turunan {spec['name']}", 2)
            add_body(doc, f"Role: {spec['role']}. Senjata: {WEAPON_LABELS.get(spec['weapon'], spec['weapon'])}. {spec['description']}")
            add_heading(doc, f"Skill {spec['name']}", 3)
            add_skill_table(doc, special_skills[spec_key])
            add_heading(doc, f"Passive {passive['name']}", 3)
            add_passive_table(doc, passive, DATA["passiveEffects"].get(passive["id"], {}))
            capstone = {"name": f"Warisan {spec['name']}", "description": f"Puncak latihan {spec['name']}: memperkuat passive utama. Memerlukan Mastery dan passive utama level 3.", "maxLevel": 1, "unlockLevel": 50}
            add_heading(doc, f"Capstone Warisan {spec['name']}", 3)
            add_passive_table(doc, capstone, DATA["passiveEffects"].get(f"{spec_key}-capstone", {}), capstone=True)

    add_heading(doc, "Catatan aturan skill", 1)
    notes = [
        "Active skill memiliki rank maksimum 5. Passive spesialisasi dan passive Core Job memiliki rank maksimum 3; Capstone memiliki rank maksimum 1.",
        "Skill slot 1 sampai 3 Special Job terbuka pada level 25. Skill slot 4 Special Job terbuka pada level 45. Skill slot 4 Adventurer Nova Fajar juga terbuka pada level 45.",
        "Core Job dipilih pada level 10 setelah seluruh equipment dilepas. Proses promosi mereset skill dan passive promotion serta mengosongkan Primary Hotbar, sesuai aturan game saat ini.",
        "Special Job dipilih pada level 25 dan harus berasal dari Core Job yang sedang dipakai. Special Job memberi passive spesialisasi dan membuka empat skill jalurnya.",
        "Capstone tidak dapat dipelajari sebelum Mastery Quest selesai dan passive utama spesialisasi mencapai level 3. Bonus capstone di registry dibuat dua kali bonus passive utama.",
        "Weapon requirement hanya dicantumkan pada skill yang benar-benar menggunakannya. Skill tanpa daftar requirement dapat digunakan tanpa syarat senjata tambahan, sesuai registry.",
    ]
    for note in notes:
        paragraph = doc.add_paragraph(style="List Bullet")
        style_paragraph(paragraph, space_after=5, line=1.08)
        style_run(paragraph.add_run(note), size=10.0)

    add_heading(doc, "Ringkasan jumlah konten", 1)
    add_table(doc, ["Jenis", "Jumlah", "Rincian"], [
        ["Tahap awal", "1", "Adventurer"],
        ["Core Job", "5", "Warrior, Rogue, Hunter, Wizard, Acolyte"],
        ["Special Job", "10", "2 turunan untuk setiap Core Job"],
        ["Active skill", "64", "4 Adventurer + 20 Core Job + 40 Special Job"],
        ["Passive node", "26", "1 Adventurer + 5 foundation + 10 specialization + 10 capstone"],
    ], [1.45, 0.75, 4.8], font_size=9.0)
    add_body(doc, "Nama, urutan, syarat, dan deskripsi di atas mengikuti registry skill/job yang sedang dipakai Lumenfall. Dokumen ini adalah katalog referensi; perubahan registry berikutnya perlu dibuat ulang agar katalog tetap sinkron.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
