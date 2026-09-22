from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\USER\Documents\ChatGPT\PROJECT GAME NGEMPER")
OUTPUT = ROOT / "exports" / "Lumenfall_Rune_Stat_Effects_Reference.docx"


RARITIES = [
    ("Cracked", 0.3, "#8F9698", "Common", 1, 1),
    ("Simple", 0.5, "#C3C8C9", "Uncommon", 1, 1),
    ("Refined", 1.0, "#57BD78", "Uncommon", 1, 2),
    ("Rare", 1.5, "#4E9DF5", "Rare", 2, 2),
    ("Epic", 2.2, "#B369E8", "Epic", 2, 3),
    ("Legendary", 3.0, "#E9BD4E", "Legendary", 3, 3),
    ("Ancient", 4.0, "#FC6C3D", "Mythic", 3, 4),
]

THEMES = [
    {
        "name": "Rune of Might",
        "category": "Offensive",
        "stats": [
            ("Physical Damage", "percent", "physical"),
            ("Strength", "flat", "flat"),
            ("Critical Damage", "percent", "percent"),
            ("Damage terhadap Boss", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Precision",
        "category": "Offensive",
        "stats": [
            ("Critical Rate", "percent", "percent"),
            ("Accuracy", "percent", "percent"),
            ("Critical Damage", "percent", "percent"),
            ("Weak Point Damage", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Swiftness",
        "category": "Offensive and Mobility",
        "stats": [
            ("Attack Speed", "percent", "percent"),
            ("Movement Speed", "percent", "percent"),
            ("Cooldown Reduction", "percent", "percent"),
            ("Evasion", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Vitality",
        "category": "Defensive",
        "stats": [
            ("Max HP", "flat", "hp"),
            ("HP Recovery", "flat", "flat"),
            ("Defense", "flat", "flat"),
            ("Damage Reduction", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Arcana",
        "category": "Magic",
        "stats": [
            ("Magic Attack", "percent", "percent"),
            ("Max MP", "flat", "flat"),
            ("Skill Damage", "percent", "percent"),
            ("Cooldown Reduction", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Focus",
        "category": "Magic and Resource",
        "stats": [
            ("MP Recovery", "flat", "flat"),
            ("Critical Rate", "percent", "percent"),
            ("Mana Cost Reduction", "percent", "percent"),
            ("Accuracy", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Elements",
        "category": "Elemental and Magic",
        "stats": [
            ("Elemental Damage", "percent", "percent"),
            ("Elemental Resistance", "percent", "percent"),
            ("Skill Damage", "percent", "percent"),
            ("Magic Attack", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Fortune",
        "category": "Economy and Loot",
        "stats": [
            ("EXP Gain", "percent", "percent"),
            ("GOLD Drop Rate", "percent", "percent"),
            ("Item Drop Rate", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of the Guardian",
        "category": "Defensive",
        "stats": [
            ("Block Rate", "percent", "percent"),
            ("Parry Rate", "percent", "percent"),
            ("Resist Stun", "percent", "percent"),
            ("Damage Reduction", "percent", "percent"),
        ],
    },
    {
        "name": "Rune of Shadows",
        "category": "Offensive and Mobility",
        "stats": [
            ("Evasion", "percent", "percent"),
            ("Critical Rate", "percent", "percent"),
            ("Critical Damage", "percent", "percent"),
            ("Attack Speed", "percent", "percent"),
        ],
    },
]

SPECIAL_RUNES = [
    ("Embercore Rune", "Elements", "Twin Elemental Lord · Dataran Bara-Beku", "Tidak ada", "Serangan api memiliki peluang meninggalkan bara."),
    ("Primordial Root Rune", "Vitality", "Ancient Treant · Padang Arunika", "Tidak ada", "Pemulihan HP meningkat saat HP rendah."),
    ("Caroq Shadow Rune", "Shadows", "Field Boss Rimba Bisik", "Rogue", "Sinergi maksimum untuk Rogue dan serangan dari bayangan."),
    ("Skywarden Rune", "Guardian", "Field Boss Tambang Selubung Besi", "Warrior", "Perfect guard memperkuat pertahanan singkat."),
    ("Jayantara's Eye Rune", "Arcana", "Field Boss Reruntuhan Tenggelam", "Wizard", "Skill elemental memperoleh penetrasi ringan."),
    ("Meteor King Rune", "Might", "Meteorfall Overlord · Benteng Hujan Meteor", "Tidak ada", "Peluang mengabaikan sebagian defense boss."),
]


def js_round(value: float) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def value_range(kind: str, power: float) -> tuple[float, float]:
    if kind == "hp":
        low, high = 60, 120
    elif kind == "flat":
        low, high = 1, 4
    elif kind == "physical":
        low, high = 4.5, 6.25
    else:
        low, high = 2, 6
    return js_round(low * power), js_round(high * power)


def fmt_number(value: float) -> str:
    if value == int(value):
        return str(int(value))
    return f"{value:.1f}"


def fmt_range(kind: str, unit: str, power: float) -> str:
    low, high = value_range(kind, power)
    suffix = "%" if unit == "percent" else ""
    return f"+{fmt_number(low)}{suffix} sampai +{fmt_number(high)}{suffix}"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def set_font(run, name="Aptos", size=10, bold=False, color="000000", italic=False):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=5, line_spacing=1.08):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 12, bold=True, color="000000")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_keep_with_next(p)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    run = p.add_run(text)
    set_font(run, size=10)
    style_paragraph(p, 5, 1.08)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_font(run, size=10)
    style_paragraph(p, 2, 1.05)
    return p


def add_table(doc, headers, rows, widths=None, header_fill="24364B", font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        set_cell_margins(cell, 90, 90, 90, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_font(run, size=font_size, bold=True, color="FFFFFF")
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        fill = "FFFFFF" if row_index % 2 == 0 else "F3F6F8"
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_shading(cell, fill)
            set_cell_borders(cell)
            set_cell_margins(cell, 90, 90, 90, 90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_font(run, size=font_size, color="000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_rarity_legend(doc):
    rows = []
    for name, power, color, mapped, min_count, max_count in RARITIES:
        count = str(min_count) if min_count == max_count else f"{min_count} sampai {max_count}"
        rows.append((name, f"{power:g}x", mapped, count, color))
    table = add_table(doc, ["Rarity Rune", "Power", "Warna katalog", "Jumlah opsi", "Kode warna"], rows,
                      widths=[1.45, 0.8, 1.35, 1.1, 1.3], font_size=8.6)
    for row_index, (_, _, color, _, _, _) in enumerate(RARITIES, start=1):
        cell = table.rows[row_index].cells[0]
        set_cell_shading(cell, color)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, size=8.6, bold=True, color="FFFFFF" if color not in ["#C3C8C9", "#E9BD4E"] else "000000")
    return table


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    if "Small Note" not in [s.name for s in styles]:
        note_style = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        note_style.font.name = "Aptos"
        note_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        note_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        note_style.font.size = Pt(8.5)
        note_style.font.italic = True
        note_style.font.color.rgb = RGBColor(75, 75, 75)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LUMENFALL  RUNE REFERENCE")
    set_font(run, size=8, bold=True, color="5A6570")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Current runtime data reference")
    set_font(run, size=8, color="6A6A6A")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Referensi Efek Stat Rune Lumenfall")
    set_font(run, name="Aptos Display", size=24, bold=True, color="000000")
    title.paragraph_format.space_after = Pt(3)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Daftar kategori dan rentang nilai dari Cracked hingga Ancient")
    set_font(run, size=11, color="5A6570")
    subtitle.paragraph_format.space_after = Pt(13)

    add_body(doc, "Dokumen ini merangkum seluruh efek stat yang dapat muncul dari Socket Rune pada registry game saat ini. Daftar disusun berdasarkan tema Rune, kategori fungsi, jumlah opsi per rarity, dan rentang nilai roll aktual. Gunakan dokumen ini sebagai referensi balancing, tooltip, dan pemeriksaan hasil drop.")
    add_body(doc, "Sumber runtime: registry Rune dan fungsi rollRuneAffixes pada lib/game/items.ts. Nilai memakai pembulatan satu angka desimal seperti implementasi game. Rune Optimizer tidak termasuk dalam rentang ini karena merupakan item modifikasi equipment yang terpisah.")

    add_heading(doc, "Urutan Rarity Rune", 1)
    add_body(doc, "Semakin tinggi rarity, semakin besar pengali nilai dan semakin banyak opsi yang dapat dimiliki. Pemetaan warna katalog di bawah mengikuti rarity item yang dipakai UI game.")
    add_rarity_legend(doc)

    add_heading(doc, "Band Nilai Runtime", 1)
    add_body(doc, "Semua stat Rune memakai salah satu band dasar berikut sebelum dikalikan power rarity. Band ini menjelaskan mengapa Max HP berbeda skala dari stat flat lain, dan mengapa Physical Damage memakai rentang khusus.")
    band_rows = [
        ("Max HP", "flat", "60 sampai 120", "Max HP"),
        ("Flat non HP", "flat", "1 sampai 4", "Strength, Defense, Max MP, HP Recovery, MP Recovery"),
        ("Physical Damage", "percent", "4.5 sampai 6.25", "Physical Damage"),
        ("Persentase standar", "percent", "2 sampai 6", "Stat persentase Rune lainnya"),
    ]
    add_table(doc, ["Band dasar", "Unit", "Rentang dasar", "Stat yang menggunakannya"], band_rows,
              widths=[1.55, 1.0, 1.55, 5.7], font_size=8.8)

    add_heading(doc, "Daftar Efek Stat Berdasarkan Tema", 1)
    add_body(doc, "Setiap tabel di bawah menunjukkan semua stat dalam satu tema dan rentang roll dari rarity paling rendah ke paling tinggi. Tanda persen berarti bonus persentase; tanpa tanda persen berarti nilai flat.")

    for theme in THEMES:
        add_heading(doc, theme["name"], 2)
        add_body(doc, f"Kategori: {theme['category']}. Pool stat: {len(theme['stats'])} pilihan; satu Rune tidak selalu mengambil seluruh pool karena jumlah opsi mengikuti rarity.")
        rows = []
        for label, unit, kind in theme["stats"]:
            values = [fmt_range(kind, unit, power) for _, power, _, _, _, _ in RARITIES]
            rows.append((label, "Persen" if unit == "percent" else "Flat", *values))
        add_table(doc, ["Efek stat", "Unit"] + [r[0] for r in RARITIES], rows,
                  widths=[1.72, 0.72, 1.18, 1.18, 1.18, 1.18, 1.18, 1.18, 1.18],
                  header_fill="334E68", font_size=7.6)

    add_heading(doc, "Rune Unik Field Boss", 1)
    add_body(doc, "Rune unik Field Boss memiliki rarity item Mythic dan runeRarity Ancient. Affix numeriknya mengikuti pool tema Ancient, sedangkan unique effect di bawah adalah efek khusus tambahan yang tersimpan pada item.")
    special_rows = [(name, theme, source, job, effect) for name, theme, source, job, effect in SPECIAL_RUNES]
    add_table(doc, ["Rune unik", "Tema", "Sumber drop", "Job restriction", "Unique effect"], special_rows,
              widths=[1.7, 1.05, 2.9, 1.05, 4.0], header_fill="6B3A4A", font_size=8.1)

    add_heading(doc, "Catatan Kategori Dan Batasan", 1)
    add_bullet(doc, "Offensive: Might, Precision, Swiftness, dan Shadows berfokus pada serangan, akurasi, critical, kecepatan, atau mobilitas.")
    add_bullet(doc, "Defensive: Vitality dan Guardian berfokus pada HP, Defense, pengurangan damage, block, parry, dan resist stun.")
    add_bullet(doc, "Magic and Elemental: Arcana, Focus, dan Elements berfokus pada Magic Attack, Skill Damage, resource, cooldown, dan efek elemental.")
    add_bullet(doc, "Economy and Loot: Fortune berfokus pada EXP, GOLD, dan peluang drop item.")
    add_bullet(doc, "Rune biasa pada registry saat ini tidak memiliki pembatasan Core Job. Pembatasan job diterapkan pada Rune unik tertentu, yaitu Rogue, Warrior, dan Wizard sesuai tabel Field Boss.")
    add_bullet(doc, "Rune of Elements saat ini memakai stat generik Elemental Damage dan Elemental Resistance. Data runtime belum memisahkan Fire, Water, Wind, Earth, Light, dan Dark sebagai stat terpisah.")
    add_bullet(doc, "Satu Rune memiliki jumlah opsi sesuai tabel rarity, lalu memilih stat secara acak tanpa pengulangan dari pool temanya. Karena itu hasil aktual satu Rune dapat berbeda dari Rune lain dengan tema dan rarity sama.")

    add_heading(doc, "Ringkasan Untuk Pembacaan Cepat", 1)
    summary_rows = [
        ("Paling rendah", "Cracked", "1 opsi", "0.3x", "Nilai terendah"),
        ("Dasar", "Simple", "1 opsi", "0.5x", "Nilai rendah"),
        ("Menengah", "Refined", "1 sampai 2 opsi", "1.0x", "Nilai standar"),
        ("Tinggi", "Rare", "2 opsi", "1.5x", "Nilai tinggi"),
        ("Sangat tinggi", "Epic", "2 sampai 3 opsi", "2.2x", "Nilai premium"),
        ("Endgame", "Legendary", "3 opsi", "3.0x", "Nilai kuat"),
        ("Puncak", "Ancient", "3 sampai 4 opsi", "4.0x", "Nilai tertinggi dan Rune unik"),
    ]
    add_table(doc, ["Tingkat", "Rarity", "Jumlah opsi", "Power", "Keterangan"], summary_rows,
              widths=[1.45, 1.3, 1.45, 0.9, 4.9], header_fill="24364B", font_size=8.7)

    doc.core_properties.title = "Referensi Efek Stat Rune Lumenfall"
    doc.core_properties.subject = "Daftar kategori dan rentang nilai stat Rune"
    doc.core_properties.author = "Lumenfall"
    doc.core_properties.comments = "Generated from the current runtime Rune registry."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
