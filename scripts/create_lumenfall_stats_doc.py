from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'LUMENFALL_Struktur_Stat_Karakter.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, 'FFFFFF'); shade(t.rows[0].cells[i], '1F4E5F')
        if widths: t.rows[0].cells[i].width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if widths: cells[i].width = Inches(widths[i])
            if ri % 2 == 1: shade(cells[i], 'F2F6F7')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text); return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
for name, size, color in [('Title', 24, '000000'), ('Heading 1', 16, '000000'), ('Heading 2', 12.5, '1F4E5F')]:
    st = styles[name]; st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name); st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)

title = doc.add_paragraph(style='Title'); title.add_run('LUMENFALL Struktur Stat Karakter')
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(12); r = sub.add_run('Referensi teknis untuk atribut, stat turunan, combat, job, skill, dan Combat Power'); r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string('4F6B73')
doc.add_paragraph('Dokumen ini menjelaskan struktur statistik karakter yang digunakan LUMENFALL. Fokusnya adalah arsitektur data dan alur perhitungan karakter, bukan katalog bonus atau daftar statistik equipment. Kesimpulan utamanya: game menggunakan satu pipeline final derived stats sebagai sumber gameplay dan Combat Power.')

doc.add_heading('1 Ruang Lingkup Sistem', level=1)
doc.add_paragraph('Stat karakter dibentuk dari data dasar karakter, profile job atau specialization, atribut yang dialokasikan, passive, skill, buff, dan sumber modifier lain. Semua sumber tersebut akhirnya diproses menjadi final derived stats yang dibaca oleh combat, UI, save validation, dan Combat Power.')
table(doc, ['Lapisan', 'Isi utama', 'Peran'], [
    ('Base character', 'Level, HP dasar, attack dasar, resource, posisi dan progresi', 'Menyediakan state karakter yang disimpan'),
    ('Primary attributes', 'STR, VIT, DEX, INT', 'Input utama yang mengubah stat turunan'),
    ('Job profile', 'Base HP, growth, attack, cooldown, range, weapon dan resource', 'Memberi identitas combat job'),
    ('Modifiers', 'Passive, buff, pet, rune, socket dan sumber loadout lain', 'Menambah atau mengubah nilai final'),
    ('Derived stats', 'HP, attack, defense, critical, speed, recovery dan utility', 'Dipakai langsung oleh gameplay'),
    ('Combat evaluation', 'DPS, EHP, sustain, utility dan special effects', 'Menghasilkan Combat Power'),
], [1.25, 3.25, 2.1])

doc.add_heading('2 Atribut Utama', level=1)
doc.add_paragraph('Atribut utama disimpan pada allocatedStats. VIT adalah nama kanonik; STA tetap diterima sebagai alias read-only untuk kompatibilitas save lama.')
table(doc, ['Atribut', 'Dampak utama pada final stats'], [
    ('STR', 'Physical Attack melalui skala serangan fisik'),
    ('VIT', 'Max HP, Physical Defense, Tenacity, HP Regen dan Max Stamina'),
    ('DEX', 'Accuracy, Evasion, Critical Rate dan Attack Speed'),
    ('INT', 'Magic Attack, Magic Defense, Max MP, Healing Power dan Skill Power'),
], [1.3, 5.3])

doc.add_heading('3 Resource dan State Pertempuran', level=1)
bullet(doc, 'HP dan Max HP: kondisi hidup karakter dan batas HP efektif.')
bullet(doc, 'Mana atau MP dan Max MP: resource skill.')
bullet(doc, 'Stamina dan Max Stamina: resource pergerakan atau aksi yang dikonfigurasi game.')
bullet(doc, 'Barrier: nilai pelindung sementara.')
bullet(doc, 'Active Buffs dan Status Effects: modifier sementara yang dibaca oleh gameplay.')
bullet(doc, 'HP Recovery dan Mana Recovery: pemulihan resource dari waktu ke waktu.')
doc.add_paragraph('Rumus dasar yang digunakan:')
table(doc, ['Nilai', 'Struktur perhitungan'], [
    ('Max HP', 'Job Base HP + level growth + VIT × 10 + modifier HP'),
    ('Max MP', '100 + INT × 6 + resource modifier'),
    ('Max Stamina', '100 + VIT × 4 + stamina modifier'),
    ('Barrier', '35 + Max HP × 0.18 saat barrier dibuat'),
], [1.5, 5.1])

doc.add_heading('4 Derived Stats Karakter', level=1)
doc.add_paragraph('DerivedStats adalah struktur runtime utama. Nilainya dihitung melalui derivedStats(hero), yang juga diekspor sebagai calculateFinalCharacterStats(hero).')
table(doc, ['Kelompok', 'Stat yang termasuk'], [
    ('Resource', 'maxHP, maxMana, staminaMax'),
    ('Offense', 'physicalAttack, attack, magicAttack, criticalRate, criticalDamage, attackSpeed'),
    ('Defense', 'physicalDefense, defense, magicDefense, damageReduction, tenacity, staggerResistance'),
    ('Combat utility', 'accuracy, evasion, blockRate, movementSpeed, cooldownReduction'),
    ('Skill and resource', 'healingPower, skillPower, skillDamage, manaCostReduction, manaRecovery, hpRecovery'),
    ('Penetration and progression', 'physicalPenetration, magicPenetration, bossDamage, eliteDamage, expGain, goldDropRate, itemDropRate, materialDropRate'),
], [1.65, 4.95])

doc.add_heading('5 Formula Stat Turunan Utama', level=1)
table(doc, ['Stat', 'Formula inti'], [
    ('Physical Attack', '(Job Attack + weapon contribution + STR × 2 + attack modifier) × attack multiplier'),
    ('Magic Attack', 'Job Attack + attack modifier + magic modifier + INT × 2'),
    ('Physical Defense', '8 + level × 1.5 + VIT × 0.5 + defense modifier'),
    ('Magic Defense', '8 + level × 1.5 + INT × 0.5 + magic defense modifier'),
    ('Accuracy', '90 + DEX + accuracy modifier'),
    ('Critical Rate', 'min(100, 2 + DEX × 0.1 + critical modifier)'),
    ('Critical Damage', '150 + critical damage modifier'),
    ('Attack Speed', '100 + DEX × 0.15 + attack speed modifier'),
    ('Evasion', 'DEX × 0.1 + evasion modifier'),
    ('Healing Power', 'INT × 0.25 + healing modifier'),
    ('Skill Power', 'INT × 0.45 + skill power modifier'),
], [1.7, 4.9])

doc.add_heading('6 Job dan Specialization', level=1)
doc.add_paragraph('Job tidak mengubah formula derived stats secara terpisah-pisah. Sistem memilih satu combat profile berdasarkan urutan specialization, core job, legacy job, atau Adventurer. Profile tersebut menyediakan base HP, HP growth, base attack, attack growth, cooldown, range, weapon type, warna dan nama resource.')
doc.add_paragraph('Karena profile dibaca secara terpusat, perubahan job otomatis memicu perhitungan ulang HP, attack, cooldown, skill yang tersedia, weapon requirement, dan Combat Power.')

doc.add_heading('7 Skill dan Passive', level=1)
bullet(doc, 'Skill damage memakai skillCombatScaling untuk menentukan physical coefficient, magic coefficient dan damage type.')
bullet(doc, 'Skill damage menggabungkan base damage, physical atau magic attack, level skill, mastery, skill power dan skill damage.')
bullet(doc, 'Skill healing memakai base heal, Magic Attack, INT dan Healing Power.')
bullet(doc, 'Passive yang aktif dipilih berdasarkan job, specialization, level, tier dan mastery requirement.')
bullet(doc, 'Skill tidak hanya bergantung pada job; weapon requirement, level, skill level, cooldown dan mana juga divalidasi.')

doc.add_heading('8 Combat Mechanics', level=1)
table(doc, ['Mekanik', 'Aturan runtime'], [
    ('Mitigation', 'Defense efektif dibatasi minimal nol setelah penetration; reduction = defense / (defense + 500 + attacker level × 10)'),
    ('Critical', 'Chance dibatasi maksimum 80 persen'),
    ('Evasion', 'Chance dibatasi maksimum 50 persen'),
    ('Block', 'Chance dibatasi maksimum 50 persen; block mengurangi damage sebesar 30 persen'),
    ('Damage Reduction', 'Memakai soft cap dan buff aktif; tidak dihitung sebagai stat mentah tanpa batas'),
], [1.55, 5.05])

doc.add_heading('9 Alur Combat Power', level=1)
doc.add_paragraph('Combat Power bukan gear score dan bukan angka tetap berdasarkan rarity. Calculator membaca final derived stats serta action kit yang benar-benar bisa digunakan karakter.')
doc.add_paragraph('Alurnya:')
for text in ['Hero state dan job profile', 'Final derived stats', 'Skill dan basic attack yang valid', 'Expected damage terhadap target benchmark', 'Physical dan magic survivability', 'Sustain, utility dan special effects', 'Total Combat Power dan breakdown UI']:
    bullet(doc, text)
table(doc, ['Komponen', 'Makna'], [
    ('Offensive Power', 'Expected physical dan magic DPS, critical, attack speed dan damage skill'),
    ('Defensive Power', 'Physical EHP dan Magic EHP dari Max HP, defense, mitigation, evasion, block dan damage reduction'),
    ('Sustain Power', 'Healing per second dan shield per second yang benar-benar dapat dihasilkan'),
    ('Utility Power', 'Kontribusi sekunder seperti Move Speed dengan cap terhadap core power'),
    ('Special Effect Power', 'Hanya efek khusus yang memiliki evaluator runtime; deskripsi pasif tidak otomatis diberi nilai'),
], [1.7, 4.9])
doc.add_paragraph('Combat Power dihitung sebagai jumlah komponen tersebut dikalikan display scale terpusat, lalu dibulatkan. Nilai dapat di-cache, tetapi cache bukan source of truth; perubahan dependency akan menyebabkan kalkulasi ulang.')

doc.add_heading('10 Data Karakter yang Disimpan', level=1)
table(doc, ['Area data', 'Field utama'], [
    ('Identitas', 'characterId, characterName, gender, appearance'),
    ('Progresi', 'level, xp, gold, kills, job, jobTier, coreJob, specialization'),
    ('Atribut', 'allocatedStats, statPoints'),
    ('Skill', 'skillLevels, skillPoints, passiveLevels, masteryChoices'),
    ('Resource runtime', 'hp, mana, stamina, barrier, activeBuffs, statusEffects'),
    ('World state', 'currentCity, currentField, inCity, posisi, quest dan progression records'),
], [1.65, 4.95])
doc.add_paragraph('Derived stats seperti Physical Attack, Magic Defense, Critical Rate dan Combat Power tidak perlu menjadi sumber permanen di save. Nilai tersebut dihitung kembali dari state karakter ketika digunakan.')

doc.add_heading('11 Kesimpulan Arsitektur', level=1)
doc.add_paragraph('Struktur statistik Lumenfall bersifat derived dan terpusat. Atribut, job, skill, passive, buff dan modifier lain mengalir ke satu perhitungan final. Gameplay dan Combat Power membaca hasil yang sama sehingga perubahan build karakter tercermin pada combat dan UI secara konsisten.')
doc.add_paragraph('File kode utama yang menjadi referensi implementasi adalah lib/game/rules.ts, lib/game/combat-mechanics.ts, lib/game/skills.ts, lib/game/combat-power.ts dan lib/game/character-screen.ts.')

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('LUMENFALL · Struktur Stat Karakter'); fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string('6B7D82')
doc.core_properties.title = 'LUMENFALL Struktur Stat Karakter'
doc.core_properties.subject = 'Struktur atribut dan statistik karakter Lumenfall'
doc.core_properties.author = 'LUMENFALL'
doc.save(OUT)
print(OUT)
